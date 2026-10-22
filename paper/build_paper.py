"""Build Context_Rot_Research_Final.docx on top of template.docx.

The template (paper/template/ieee_template.docx, a text-free copy of an IEEE-style two-column conference paper) is loaded as the base
document. Its page setup, section layout (single-column title, multi-column author block, two-column
body), styles (Heading 1-3, Caption, List Paragraph, Table Grid), numbering definitions (roman section
numbers, lettered subsections, bullets, [n] references) and run-level formatting conventions
(Times New Roman 10 pt body, 9 pt authors/references, 8 pt tables) are reused. Only the body content
is replaced, and the template's old images and links are dropped.

Run from the repository root:  python paper/build_paper.py
"""
import copy
import datetime as dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paper_content import ABSTRACT, AUTHORS, BODY, KEYWORDS, REFERENCES, TITLE  # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = HERE / "template" / "ieee_template.docx"
FIGURES = HERE / "figures"
OUT = HERE / "Context_Rot_Research_Final.docx"

TNR = "Times New Roman"
MONO = "Courier New"
BODY_PT = 10
COLUMN_TWIPS = 4590            # (12240 - 2 * 1440 - 180) / 2: one body column
BULLET_NUM_ID = 15             # template bullet list (Symbol bullet, 0.25" hanging)
REFERENCE_NUM_ID = 12          # template "[%1]" reference list

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]


# --------------------------------------------------------------------------- low-level XML helpers
def mk(tag, attrs=None, children=()):
    e = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        e.set(qn(k), str(v))
    for c in children:
        e.append(c)
    return e


def local(tag):
    return tag.split("}")[-1]


def ppr_insert(p, child):
    """Insert a pPr child at its schema position, replacing any existing element of that kind."""
    pPr = p._p.get_or_add_pPr()
    name = local(child.tag)
    for old in pPr.findall(child.tag):
        pPr.remove(old)
    rank = PPR_ORDER.index(name)
    for existing in pPr:
        ename = local(existing.tag)
        if ename in PPR_ORDER and PPR_ORDER.index(ename) > rank:
            existing.addprevious(child)
            return child
    pPr.append(child)
    return child


def mark_rpr(size_pt, font=TNR, bold=False, italic=False):
    """Paragraph-mark run properties, so empty lines and line heights follow the template."""
    children = [mk("w:rFonts", {"w:ascii": font, "w:hAnsi": font, "w:cs": font})]
    if bold:
        children.append(mk("w:b"))
    if italic:
        children.append(mk("w:i"))
    half = int(round(size_pt * 2))
    children += [mk("w:sz", {"w:val": half}), mk("w:szCs", {"w:val": half})]
    return mk("w:rPr", children=children)


def fmt_run(run, size=BODY_PT, font=TNR, bold=None, italic=None, color=None, valign=None):
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    rPr.rFonts.set(qn("w:cs"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if valign == "super":
        run.font.superscript = True
    elif valign == "sub":
        run.font.subscript = True
    return run


def numbered(p, num_id, ilvl=0):
    ppr_insert(p, mk("w:numPr", children=[mk("w:ilvl", {"w:val": ilvl}), mk("w:numId", {"w:val": num_id})]))


# --------------------------------------------------------------------------- text markup
TAG_RE = re.compile(r"<(b|i|code|super|sub)>(.*?)</\1>", re.S)


class Refs:
    """IEEE numbering: references are numbered in order of first citation."""

    def __init__(self, entries):
        self.entries = entries
        self.order = []

    def number(self, key):
        key = key.strip()
        if key not in self.entries:
            raise KeyError(f"unknown reference key: {key}")
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def cite(self, keys):
        nums = sorted({self.number(k) for k in keys.split(",")})
        groups, run = [], [nums[0]]
        for n in nums[1:]:
            if n == run[-1] + 1:
                run.append(n)
            else:
                groups.append(run)
                run = [n]
        groups.append(run)
        parts = []
        for g in groups:
            if len(g) >= 3:
                parts.append(f"[{g[0]}]–[{g[-1]}]")
            else:
                parts.extend(f"[{n}]" for n in g)
        return ", ".join(parts)


def resolve(text, labels, refs, prose=True):
    if prose:
        text = " ".join(text.split())
    text = re.sub(r"\[#([\w:-]+)\]", lambda m: str(labels[m.group(1)]), text)
    if refs is not None:
        text = re.sub(r"\[@([^\]]+)\]", lambda m: refs.cite(m.group(1)), text)
    return text


def add_rich(p, text, size=BODY_PT, bold=None, italic=None, color=None):
    """Add runs for text containing <b>/<i>/<code>/<super>/<sub> markup."""
    pos = 0
    for m in TAG_RE.finditer(text):
        if m.start() > pos:
            fmt_run(p.add_run(text[pos:m.start()]), size, bold=bold, italic=italic, color=color)
        tag, inner = m.group(1), m.group(2)
        if tag == "b":
            add_rich(p, inner, size, bold=True, italic=italic, color=color)
        elif tag == "i":
            add_rich(p, inner, size, bold=bold, italic=not italic if italic else True, color=color)
        elif tag == "code":
            fmt_run(p.add_run(inner), size - 1, font=MONO, bold=bold, color=color)
        else:
            fmt_run(p.add_run(inner), size, bold=bold, italic=italic, color=color, valign=tag)
        pos = m.end()
    if pos < len(text):
        fmt_run(p.add_run(text[pos:]), size, bold=bold, italic=italic, color=color)


def add_field(p, instr, cached, **fmt):
    fmt_run(p.add_run(), **fmt)._r.append(mk("w:fldChar", {"w:fldCharType": "begin"}))
    r = fmt_run(p.add_run(), **fmt)
    it = mk("w:instrText", {"xml:space": "preserve"})
    it.text = f" {instr} "
    r._r.append(it)
    fmt_run(p.add_run(), **fmt)._r.append(mk("w:fldChar", {"w:fldCharType": "separate"}))
    fmt_run(p.add_run(cached), **fmt)
    fmt_run(p.add_run(), **fmt)._r.append(mk("w:fldChar", {"w:fldCharType": "end"}))


# --------------------------------------------------------------------------- builder
class PaperBuilder:
    def __init__(self, template):
        self.doc = Document(template)
        self.body = self.doc.element.body
        self._capture_template_parts()
        self._clear_body()
        self._add_step_numbering()
        self.labels = {}
        self.refs = Refs(REFERENCES)

    # ---- template capture -------------------------------------------------------------------
    def _capture_template_parts(self):
        paras = self.body.findall(qn("w:p"))

        def sect_of(p_el):
            return p_el.find(qn("w:pPr")).find(qn("w:sectPr"))

        # Title paragraph (ends the first, single-column section) and the spacer section after it.
        self.title_p = copy.deepcopy(paras[0])
        self.spacer_ps = [copy.deepcopy(paras[1]), copy.deepcopy(paras[2])]
        for el in self.title_p.findall(qn("w:bookmarkStart")) + self.title_p.findall(qn("w:bookmarkEnd")):
            self.title_p.remove(el)
        # Author block section (three columns in the template; four authors here, so two columns).
        self.author_sect = copy.deepcopy(sect_of(paras[20]))
        cols = self.author_sect.find(qn("w:cols"))
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "720")
        # Two-column body section and the single-column section used for full-width figures/tables.
        self.two_col_sect = copy.deepcopy(sect_of(paras[21]))
        self.one_col_sect = copy.deepcopy(sect_of(paras[67]))
        for s in (self.two_col_sect, self.one_col_sect):
            assert s.find(qn("w:type")).get(qn("w:val")) == "continuous"

    def _clear_body(self):
        final = self.body.find(qn("w:sectPr"))
        for child in list(self.body):
            if child is not final:
                self.body.remove(child)
        # The template's last section is an A4 page holding one empty paragraph; end on the
        # two-column US Letter body layout used everywhere else.
        new_final = copy.deepcopy(self.two_col_sect)
        final.addprevious(new_final)
        self.body.remove(final)
        # Drop template images and hyperlinks that the cleared body no longer references.
        part = self.doc.part
        xml = self.body.xml
        for rId, rel in list(part.rels.items()):
            if rel.reltype in (RT.IMAGE, RT.HYPERLINK) and f'"{rId}"' not in xml:
                part.drop_rel(rId)

    def _add_step_numbering(self):
        numbering = self.doc.part.numbering_part.element
        abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        self.step_abstract = max(abstract_ids) + 1
        abstract = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{self.step_abstract}">'
            '<w:multiLevelType w:val="singleLevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1)"/>'
            '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="360" w:hanging="360"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{TNR}" w:hAnsi="{TNR}" w:cs="{TNR}"/></w:rPr></w:lvl>'
            '</w:abstractNum>')
        numbering.find(qn("w:num")).addprevious(abstract)

    def _new_step_list(self):
        numbering = self.doc.part.numbering_part.element
        nums = numbering.findall(qn("w:num"))
        num_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
        num = parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{self.step_abstract}"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>')
        nums[-1].addnext(num)
        return num_id

    # ---- generic paragraph helpers --------------------------------------------------------------
    def _append(self, el):
        self.body.find(qn("w:sectPr")).addprevious(el)

    def para(self, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=BODY_PT):
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        ppr_insert(p, mark_rpr(size))
        return p

    def text(self, text, prose=True):
        return resolve(text, self.labels, self.refs, prose)

    def section_break(self, sect_template):
        """Empty, tiny paragraph that ends the current section with the given layout."""
        p = self.para(align=None, size=2)
        p.paragraph_format.space_after = Pt(0)
        ppr_insert(p, copy.deepcopy(sect_template))
        return p

    # ---- front matter ---------------------------------------------------------------------------
    def front_matter(self):
        title = copy.deepcopy(self.title_p)
        runs = title.findall(qn("w:r"))
        for r in runs[1:]:
            title.remove(r)
        runs[0].find(qn("w:t")).text = TITLE
        self._append(title)
        for sp in self.spacer_ps:
            self._append(copy.deepcopy(sp))

        def author_line(text, column_break=False):
            p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
            if column_break:
                p.add_run().add_break(WD_BREAK.COLUMN)
            fmt_run(p.add_run(text), 9)
            return p

        # Two columns: authors 1 and 3 in the left column, 2 and 4 in the right, giving two rows.
        order = [AUTHORS[0], None, AUTHORS[2], "break", AUTHORS[1], None, AUTHORS[3]]
        pending_break = False
        last = None
        for item in order:
            if item == "break":
                pending_break = True
                continue
            if item is None:
                last = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
                continue
            name, affil, email = item
            last = author_line(name, column_break=pending_break)
            pending_break = False
            for line in affil:
                last = author_line(line)
            last = author_line(email)
        ppr_insert(last, copy.deepcopy(self.author_sect))
        self.para(size=9)  # spacer before the abstract, in the two-column body

        for i, (label, body) in enumerate(ABSTRACT):
            p = self.para()
            fmt_run(p.add_run(label), bold=True)
            fmt_run(p.add_run(" "))
            add_rich(p, self.text(body))
        kw = self.para(style="List Paragraph")
        kw.paragraph_format.left_indent = Inches(0)
        fmt_run(kw.add_run("Keywords:"), bold=True, italic=True)
        fmt_run(kw.add_run(" " + KEYWORDS + "."), italic=True)

    # ---- body blocks ------------------------------------------------------------------------------
    def assign_labels(self):
        counters = {}
        prefix_of = {"figure": "fig", "table": "tab", "eq": "eq", "code": "lst"}
        for block in BODY:
            kind = block[0]
            if kind in prefix_of:
                label = block[1]
                assert label.startswith(prefix_of[kind] + ":"), label
                counters[kind] = counters.get(kind, 0) + 1
                self.labels[label] = counters[kind]

    def heading(self, level, text):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        p.add_run(text)
        return p

    def body_par(self, text):
        p = self.para()
        add_rich(p, self.text(text))
        return p

    def bullets(self, items, num_id=BULLET_NUM_ID):
        for item in items:
            p = self.para(style="List Paragraph")
            numbered(p, num_id)
            add_rich(p, self.text(item))

    def steps(self, items):
        self.bullets(items, num_id=self._new_step_list())

    def equation(self, label, expr):
        p = self.para(align=None)
        pf = p.paragraph_format
        pf.space_before = Pt(4)
        pf.space_after = Pt(4)
        pf.keep_together = True
        ppr_insert(p, mk("w:tabs", children=[
            mk("w:tab", {"w:val": "center", "w:pos": COLUMN_TWIPS // 2}),
            mk("w:tab", {"w:val": "right", "w:pos": COLUMN_TWIPS}),
        ]))
        fmt_run(p.add_run("\t"))
        add_rich(p, self.text(expr, prose=False))
        fmt_run(p.add_run(f"\t({self.labels[label]})"))

    def caption(self, kind, label, text, keep_next=False):
        p = self.para(style="Caption", align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        pf = p.paragraph_format
        pf.keep_with_next = keep_next
        pf.keep_together = True  # never split a caption across columns or pages
        if kind != "Figure":
            pf.space_after = Pt(2)
        style = dict(size=None, italic=False, color=RGBColor(0, 0, 0))
        fmt_run(p.add_run(f"{kind} "), **style)
        add_field(p, f"SEQ {kind} \\* ARABIC", str(self.labels[label]), **style)
        fmt_run(p.add_run(": "), **style)
        add_rich(p, self.text(text), size=9, italic=False, color=RGBColor(0, 0, 0))
        return p

    def figure(self, label, filename, caption, width_in):
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        p.add_run().add_picture(str(FIGURES / filename), width=Inches(width_in))
        self.caption("Figure", label, caption)

    def code(self, label, caption, source):
        self.caption("Listing", label, caption, keep_next=True)
        lines = source.split("\n")
        too_long = [ln for ln in lines if len(ln) > 49]
        if too_long:
            raise ValueError(f"listing {label} has lines wider than one column: {too_long}")
        for i, line in enumerate(lines):
            p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, size=7.5)
            ppr_insert(p, mark_rpr(7.5, font=MONO))
            pf = p.paragraph_format
            pf.keep_with_next = i < len(lines) - 1
            pf.keep_together = True
            pf.left_indent = Pt(3)
            if i == len(lines) - 1:
                pf.space_after = Pt(6)
            ppr_insert(p, mk("w:shd", {"w:val": "clear", "w:color": "auto", "w:fill": "F2F2F2"}))
            if line:
                fmt_run(p.add_run(line), 7.5, font=MONO)

    def table(self, label, caption, header, rows, widths_in, opts):
        size = opts.get("size", 8)
        self.caption("Table", label, caption, keep_next=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        total = sum(int(w * 1440) for w in widths_in)
        tblW = tblPr.find(qn("w:tblW"))
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(total))
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_in):
            gc.set(qn("w:w"), str(int(w * 1440)))
        for r_i, values in enumerate([header] + rows):
            row = t.rows[r_i]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(mk("w:cantSplit"))
            if r_i == 0:
                trPr.append(mk("w:tblHeader"))
            for c_i, value in enumerate(values):
                cell = row.cells[c_i]
                cell.width = Inches(widths_in[c_i])
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                # Keep each table on one page/column with its caption, unless it may split (header repeats).
                p.paragraph_format.keep_with_next = r_i == 0 or (r_i < len(rows) and not opts.get("split_ok"))
                ppr_insert(p, mark_rpr(size))
                add_rich(p, self.text(value), size=size, bold=True if r_i == 0 else None,
                         color=RGBColor(0, 0, 0))
        spacer = self.para(size=6)
        spacer.paragraph_format.space_after = Pt(2)

    def references(self):
        p = self.para()
        p.paragraph_format.space_before = Pt(6)
        fmt_run(p.add_run("References:"), bold=True)
        unused = sorted(set(REFERENCES) - set(self.refs.order))
        if unused:
            raise ValueError(f"references never cited: {unused}")
        for key in self.refs.order:
            r = self.para(style="List Paragraph", size=9)
            numbered(r, REFERENCE_NUM_ID)
            # Zero-width spaces after "/" give Word places to wrap long URLs, so justified
            # reference lines are not stretched; they are invisible in print.
            def breakable(m):
                return m.group(0).replace("/", "/\u200b").rstrip("\u200b")

            entry = re.sub(r"https?://\S+", breakable, REFERENCES[key])
            add_rich(r, entry, size=9)
        # A continuous section break balances the two reference columns on the last page.
        self.section_break(self.two_col_sect)
        self.para(size=2)

    def body_blocks(self):
        self.assign_labels()
        for block in BODY:
            kind, args = block[0], block[1:]
            if kind in ("h1", "h2", "h3"):
                self.heading(int(kind[1]), args[0])
            elif kind == "p":
                self.body_par(args[0])
            elif kind == "bullets":
                self.bullets(args[0])
            elif kind == "steps":
                self.steps(args[0])
            elif kind == "eq":
                self.equation(*args)
            elif kind == "figure":
                self.figure(*args)
            elif kind == "code":
                self.code(*args)
            elif kind == "table":
                self.table(*args)
            elif kind == "full_start":
                self.section_break(self.two_col_sect)
            elif kind == "full_end":
                self.section_break(self.one_col_sect)
            elif kind == "refs":
                self.references()
            else:
                raise ValueError(kind)

    def properties(self):
        cp = self.doc.core_properties
        cp.title = TITLE
        cp.author = ", ".join(a[0] for a in AUTHORS)
        cp.last_modified_by = AUTHORS[0][0]
        cp.subject = "Measuring and repairing context rot in large language models"
        cp.keywords = KEYWORDS
        cp.comments = "Expanded research paper built on template.docx"
        cp.category = "Research paper"
        now = dt.datetime.now().replace(microsecond=0)
        cp.created = now
        cp.modified = now
        cp.revision = 1
        # The template's settings carry an editor plugin's per-document variables; they are not layout.
        doc_vars = self.doc.settings.element.find(qn("w:docVars"))
        if doc_vars is not None:
            self.doc.settings.element.remove(doc_vars)
        # Extended/custom properties still describe the template (its page and word counts, an
        # editor-plugin document id). Drop those values; Word recomputes the statistics on save.
        for part in self.doc.part.package.iter_parts():
            name = str(part.partname)
            if name == "/docProps/app.xml":
                xml = part.blob.decode("utf-8")
                for tag in ("TotalTime", "Pages", "Words", "Characters", "Lines", "Paragraphs",
                            "CharactersWithSpaces"):
                    xml = re.sub(rf"<{tag}>[^<]*</{tag}>", "", xml)
                part._blob = xml.encode("utf-8")
            elif name == "/docProps/custom.xml":
                xml = part.blob.decode("utf-8")
                xml = re.sub(r'<property [^>]*name="GrammarlyDocumentId".*?</property>', "", xml)
                part._blob = xml.encode("utf-8")

    def build(self, out):
        self.front_matter()
        self.body_blocks()
        self.properties()
        self.doc.save(out)
        return out


if __name__ == "__main__":
    path = PaperBuilder(TEMPLATE).build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT)
    print("saved", path)
