"""Dump a .docx's sections, headers/footers, paragraph styles and tables (step 1: inspection).

Usage:  python paper/inspect_docs.py <file.docx> [short]
"""
import sys
from docx import Document
from docx.shared import Emu

sys.stdout.reconfigure(encoding="utf-8")


def dump(path, full_text=True):
    d = Document(path)
    print("=" * 100)
    print("FILE:", path)
    for i, s in enumerate(d.sections):
        print(f"-- section {i}: page {s.page_width.inches if s.page_width else None}x{s.page_height.inches if s.page_height else None} in, "
              f"margins T{s.top_margin.inches if s.top_margin else None} B{s.bottom_margin.inches if s.bottom_margin else None} "
              f"L{s.left_margin.inches if s.left_margin else None} R{s.right_margin.inches if s.right_margin else None}, start={s.start_type}, "
              f"diff_first={s.different_first_page_header_footer}")
        sectPr = s._sectPr
        cols = sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols')
        if cols is not None:
            print("   cols:", cols.attrib)
        for name, hf in [("header", s.header), ("footer", s.footer), ("first_header", s.first_page_header), ("first_footer", s.first_page_footer)]:
            try:
                txt = " | ".join(p.text for p in hf.paragraphs if p.text.strip())
                print(f"   {name}: linked={hf.is_linked_to_previous} text={txt!r}")
            except Exception as e:
                print("  ", name, "err", e)
    print("-- body paragraphs:")
    body = d.element.body
    pi = 0
    ti = 0
    for child in body.iterchildren():
        tag = child.tag.split('}')[1]
        if tag == 'p':
            p = d.paragraphs[pi]
            pi += 1
            runs = p.runs
            fmt = ""
            if runs:
                r = runs[0]
                fmt = f"font={r.font.name} size={r.font.size.pt if r.font.size else None} b={r.bold} i={r.italic}"
            has_draw = 'drawing' in p._p.xml or 'pict' in p._p.xml
            sect = 'sectPr' in p._p.xml
            txt = p.text if full_text else p.text[:150]
            print(f"[P{pi-1}] style={p.style.name!r} align={p.alignment} {fmt}{' DRAWING' if has_draw else ''}{' SECTPR' if sect else ''} :: {txt}")
        elif tag == 'tbl':
            t = d.tables[ti]
            ti += 1
            print(f"[TABLE {ti-1}] style={t.style.name if t.style else None} rows={len(t.rows)} cols={len(t.columns)}")
            for r in t.rows:
                print("    | " + " | ".join(c.text.replace('\n', ' / ')[:60] for c in r.cells))
        else:
            print("[", tag, "]")


dump(sys.argv[1], full_text=(len(sys.argv) < 3))
