"""Sanity checks for Context_Rot_Research_Final.docx.

Checks that the file opens, that nothing from the template's original paper survives, that every
section heading and reference is present, that citations and captions are consistent, and that no
build markup leaked into the text. Run from the repository root:  python paper/verify_output.py
"""
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paper_content import BODY, REFERENCES  # noqa: E402

HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Context_Rot_Research_Final.docx"
TEMPLATE_LEFTOVERS = ["breast", "cancer", "CNN", "Deepak", "Abhiraj", "Malhotra", "mammogra", "infrared",
                      "Kaggle", "kaggle", "Rajpura", "Himachal", "thermal", "RFID", "Grammarly"]

failures = []


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        failures.append(msg)


doc = Document(OUT)
paras = doc.paragraphs
text_all = "\n".join(p.text for p in paras)
for t in doc.tables:
    for row in t.rows:
        text_all += "\n" + " | ".join(c.text for c in row.cells)

with zipfile.ZipFile(OUT) as z:
    names = z.namelist()
    xml_all = "".join(z.read(n).decode("utf-8", "ignore") for n in names if n.endswith((".xml", ".rels")))
    check(z.testzip() is None, "zip container is intact")
    media = [n for n in names if n.startswith("word/media/")]

leftovers = [w for w in TEMPLATE_LEFTOVERS if w in xml_all]
check(not leftovers, f"no template leftovers in any XML part {leftovers or ''}")
check(not re.search(r"\[@|\[#|</?(b|i|code|super|sub)>", text_all), "no build markup leaked into text")
n_figures = sum(1 for b in BODY if b[0] == "figure")
check(len(media) == n_figures, f"exactly the {n_figures} new figures are embedded ({len(media)} media files)")

headings = [(int(b[0][1]), b[1]) for b in BODY if b[0] in ("h1", "h2", "h3")]
doc_headings = [(int(p.style.name[-1]), p.text) for p in paras if p.style.name.startswith("Heading")]
check(doc_headings == headings, f"all {len(headings)} headings present in order")

ref_paras = [p for p in paras if p._p.pPr is not None and p._p.pPr.numPr is not None
             and p._p.pPr.numPr.numId.val == 12]
check(len(ref_paras) == len(REFERENCES), f"{len(ref_paras)} of {len(REFERENCES)} references listed")


def is_code(p):
    return any(r.font.name == "Courier New" for r in p.runs)


body_text = "\n".join(p.text for p in paras if p not in ref_paras and not is_code(p))
for t in doc.tables:
    for row in t.rows:
        body_text += "\n" + " ".join(c.text for c in row.cells)
cited = {int(n) for n in re.findall(r"\[(\d+)\]", body_text)}
for a, b in re.findall(r"\[(\d+)\]–\[(\d+)\]", body_text):
    cited |= set(range(int(a), int(b) + 1))
check(cited == set(range(1, len(REFERENCES) + 1)), f"every reference 1..{len(REFERENCES)} is cited in the text")

first_seen = []
for n in re.findall(r"\[(\d+)\]", body_text):
    if int(n) not in first_seen:
        first_seen.append(int(n))
check(first_seen == sorted(first_seen), "references are numbered in order of first citation (IEEE)")

for kind in ("Figure", "Table", "Listing"):
    nums = [int(m) for m in re.findall(rf"^{kind} (\d+):", text_all, re.M)]
    check(nums == list(range(1, len(nums) + 1)), f"{kind} captions numbered 1..{len(nums)}")

sections = doc.sections
layouts = []
for s in sections:
    cols = s._sectPr.find(qn("w:cols"))
    layouts.append(int(cols.get(qn("w:num"), "1")))
check(all(abs(s.page_width.inches - 8.5) < 0.01 and abs(s.page_height.inches - 11) < 0.01 for s in sections),
      "every section is US Letter, as in the template")
check(all(abs(m - 1.0) < 0.01 for s in sections
          for m in (s.top_margin.inches, s.bottom_margin.inches, s.left_margin.inches, s.right_margin.inches)),
      "1-inch margins everywhere, as in the template")
print(f"INFO  {len(sections)} sections, column layout {layouts}")
print(f"INFO  {len(paras)} paragraphs, {len(doc.tables)} tables, ~{len(text_all.split())} words")

sys.exit(1 if failures else 0)
